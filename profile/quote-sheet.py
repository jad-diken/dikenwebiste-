# Builds Diken-Price-Offer-Template.xlsx: the price offer as a calculating spreadsheet.
# Run: python3 quote-sheet.py
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
import os

HERE = os.path.dirname(os.path.abspath(__file__))
INK, RED, REDF, MUTED, LINE, PALE, PH = "0B0B0C", "ED1C26", "D4141D", "6B6E73", "E3E5E8", "FAFAFA", "9A9DA3"
F = "Arial"
def font(size=10, bold=False, color="111214", italic=False): return Font(name=F, size=size, bold=bold, color=color, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
side = lambda c, s="thin": Side(style=s, color=c)
hair = Border(bottom=side(LINE))
label_font = font(7.5, True, MUTED)

wb = Workbook(); ws = wb.active; ws.title = "Price Offer"
ws.sheet_view.showGridLines = False
widths = {"A": 8.5, "B": 44, "C": 8, "D": 10, "E": 14, "F": 19}
for c, w in widths.items(): ws.column_dimensions[c].width = w
COLS = "ABCDEF"

def put(ref, value, f=None, align=None, fl=None, border=None, fmt=None):
    c = ws[ref]; c.value = value
    if f: c.font = f
    if align: c.alignment = align
    if fl: c.fill = fl
    if border: c.border = border
    if fmt: c.number_format = fmt
    return c
def band(r1, r2, color, cols=COLS):
    for r in range(r1, r2 + 1):
        for c in cols: ws[f"{c}{r}"].fill = fill(color)
L, R, C = Alignment(horizontal="left", vertical="center"), Alignment(horizontal="right", vertical="center"), Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Letterhead, rows 1 to 5
for r in range(1, 6): ws.row_dimensions[r].height = 15
band(1, 5, INK); band(1, 5, REDF, "F")
logo = Image(os.path.join(HERE, "..", "assets", "brand", "diken-d-96.png")); logo.width = logo.height = 48; ws.add_image(logo, "A1")
ws.merge_cells("B1:C2"); put("B1", "DIKEN BROS", font(20, True, "FFFFFF"), Alignment(horizontal="left", vertical="bottom"))
ws.merge_cells("B3:C3"); put("B3", "POWER · PERFORMANCE · PRECISION", font(7.5, True, RED), Alignment(horizontal="left", vertical="top"))
ws.merge_cells("D2:E2"); put("D2", "EXCLUSIVE AGENT FOR JORDAN", font(6.5, True, "B8BABD"), Alignment(horizontal="left", vertical="bottom"), fill(INK)); ws.merge_cells("D3:E4")
motul = Image(os.path.join(HERE, "..", "assets", "logos", "motul-186.png")); motul.width, motul.height = 92, 25; ws.add_image(motul, "D3")
put("F2", "06 416 6660", font(9.5, True, "FFFFFF"), R); put("F3", "info@dikenbros.com", font(9, False, "FFFFFF"), R); put("F4", "dikendelivery.com", font(9, False, "FFFFFF"), R)
ws.row_dimensions[6].height = 5; band(6, 6, RED)
ws.row_dimensions[7].height = 10

# Title and meta
ws.row_dimensions[8].height = 30
ws.merge_cells("A8:B8"); put("A8", "Price Offer", font(24, True), Alignment(horizontal="left", vertical="bottom"))
ws.merge_cells("A9:B9"); put("A9", "عرض سعر", font(12, True, MUTED), Alignment(horizontal="left", vertical="center", readingOrder=2)); ws.row_dimensions[9].height = 18
ws.row_dimensions[10].height = 16
put("A10", "Distribution & Agencies", font(8, True, "FFFFFF"), C, fill(REDF)); ws.merge_cells("A10:B10")
ws["A10"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
put("C8", "OFFER", label_font, Alignment(horizontal="left", vertical="bottom")); put("E8", "YOUR REFERENCE", label_font, Alignment(horizontal="left", vertical="bottom"))
meta = [(9, "Offer no.", "DB-Q-2026-____", "RFQ no.", "[client reference]"), (10, "Date", None, "Enquiry date", None), (11, "Valid until", None, "Pages", "1 of 1"), (12, "Currency", "JOD", None, None)]
for r, k1, v1, k2, v2 in meta:
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 15)
    put(f"C{r}", k1, font(9, False, MUTED), L)
    put(f"D{r}", v1, font(9, False, PH if v1 and v1.startswith(("DB-Q", "[")) else "111214"), L)
    if k2: put(f"E{r}", k2, font(9, False, MUTED), L)
    if v2: put(f"F{r}", v2, font(9, False, PH if v2.startswith("[") else "111214"), L)
put("D10", None, font(9), L, fmt="dd mmm yyyy"); put("F10", None, font(9), L, fmt="dd mmm yyyy")
put("D11", '=IF(ISNUMBER(D10),D10+30,"")', font(9), L, fmt="dd mmm yyyy")
ws.row_dimensions[13].height = 8

# Parties
top_red = Border(top=side(RED, "medium"))
for c in "ABC": ws[f"{c}14"].border = top_red
for c in "DEF": ws[f"{c}14"].border = top_red
ws.row_dimensions[14].height = 16
put("A14", "PREPARED FOR", label_font, Alignment(horizontal="left", vertical="center")); put("D14", "PREPARED BY", label_font, Alignment(horizontal="left", vertical="center"))
party = [(15, "[Client company]", "[Your name]", True), (16, "Attn. [Name], [Title]", "Diken Bros Co., Distribution & Agencies", False), (17, "[Address, city]", "Wadi Saqra, Arar Street 14, Amman 11181", False), (18, "[email] · [phone]", "[name]@dikenbros.com · [mobile]", False)]
for r, a, d, bold in party:
    ws.row_dimensions[r].height = 16 if bold else 14
    ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}")
    put(f"A{r}", a, font(11 if bold else 9, bold, PH if a.startswith("[") or a.startswith("Attn. [") else "111214"), L)
    put(f"D{r}", d, font(11 if bold else 9, bold, PH if d.startswith("[") else "111214"), L)
ws.row_dimensions[19].height = 8

# Subject
ws.merge_cells("A20:F20"); ws.row_dimensions[20].height = 30
put("A20", "Subject: Supply of Motul lubricants for [client / fleet / site]. Prices are for genuine Motul products supplied by Diken Bros as exclusive agent for Jordan, delivered to [delivery address].", font(9.5), WRAP)
ws.row_dimensions[21].height = 8

# Items
HR = 22
ws.row_dimensions[HR].height = 18
for c, h, al in zip(COLS, ["#", "DESCRIPTION", "QTY", "UNIT", "UNIT PRICE", "TOTAL"], [C, L, R, L, R, R]):
    put(f"{c}{HR}", h, font(7.5, True, "FFFFFF"), al, fill(INK))
FIRST, N = HR + 1, 12
for i in range(N):
    r = FIRST + i
    ws.row_dimensions[r].height = 20
    shade = fill(PALE) if i % 2 else None
    put(f"A{r}", i + 1, font(9, False, MUTED), C, shade, hair)
    put(f"B{r}", "[Product, grade and pack size. Motul part number]" if i == 0 else None, font(9, False, PH if i == 0 else "111214"), Alignment(horizontal="left", vertical="center", wrap_text=True), shade, hair)
    put(f"C{r}", None, font(9), R, shade, hair, "0")
    put(f"D{r}", "case" if i == 0 else None, font(9, False, PH if i == 0 else "111214"), L, shade, hair)
    put(f"E{r}", None, font(9), R, shade, hair, "#,##0.00")
    put(f"F{r}", f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})', font(9), R, shade, hair, "#,##0.00")
LAST = FIRST + N - 1
ws.row_dimensions[LAST + 1].height = 8

# Totals
T = LAST + 2
rows = [("Subtotal", f"=SUM(F{FIRST}:F{LAST})", None), ("discount", f"=-F{T}*D{T+1}", 0.0), ("Net", f"=F{T}+F{T+1}", None), ("tax", f"=F{T+2}*D{T+3}", 0.16), ("Total, JOD", f"=F{T+2}+F{T+3}", None)]
for i, (k, fx, pct) in enumerate(rows):
    r = T + i; grand = i == 4
    ws.row_dimensions[r].height = 20 if grand else 16
    if k == "discount":
        put(f"D{r}", pct, font(9, False, MUTED), R, fmt="0%"); put(f"E{r}", "Discount", font(9), L, border=hair)
    elif k == "tax":
        put(f"D{r}", pct, font(9, False, MUTED), R, fmt="0%"); put(f"E{r}", "Sales tax", font(9), L, border=hair)
    else:
        put(f"E{r}", k, font(11 if grand else 9, grand, "FFFFFF" if grand else "111214"), L, fill(INK) if grand else None, None if grand else hair)
    put(f"F{r}", fx, font(11 if grand else 9, grand, "FFFFFF" if grand else "111214"), R, fill(INK) if grand else None, None if grand else hair, "#,##0.00")
    if grand: ws[f"E{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
G = T + 4
ws.row_dimensions[G + 1].height = 10

# Terms
S = G + 2
terms = [("VALIDITY", "30 days from the date of this offer. Prices are subject to Motul's price list at the time of order confirmation.", "PAYMENT", "As agreed by contract."),
         ("DELIVERY", "Ex-stock items within 5 working days of order confirmation. Delivery by Diken company vehicles to the address above.", "WARRANTY AND RETURNS", "Genuine Motul products with manufacturer warranty. Unopened, undamaged cases returnable within 14 days.")]
r = S
for h1, b1, h2, b2 in terms:
    ws.row_dimensions[r].height = 13
    ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}")
    put(f"A{r}", h1, label_font, Alignment(horizontal="left", vertical="bottom")); put(f"D{r}", h2, label_font, Alignment(horizontal="left", vertical="bottom"))
    r += 1; ws.row_dimensions[r].height = 28
    ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}")
    put(f"A{r}", b1, font(8.5, False, "2A2C30"), WRAP); put(f"D{r}", b2, font(8.5, False, "2A2C30"), WRAP)
    r += 1
ws.row_dimensions[r].height = 6; r += 1
ws.merge_cells(f"A{r}:F{r}"); ws.row_dimensions[r].height = 26
put(f"A{r}", "Prices in Jordanian Dinars. Sales tax at the prevailing rate. This offer is confidential and intended for the addressee. Acceptance by signature below or by purchase order referencing the offer number.", font(8, False, MUTED), WRAP)
r += 1; ws.row_dimensions[r].height = 10; r += 1

# Signatures
top_hair = Border(top=side(LINE))
for c in COLS: ws[f"{c}{r}"].border = top_hair
ws.row_dimensions[r].height = 18
ws.merge_cells(f"A{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}")
put(f"A{r}", "ACCEPTED FOR THE CLIENT", label_font, Alignment(horizontal="left", vertical="center")); put(f"D{r}", "FOR DIKEN BROS CO.", label_font, Alignment(horizontal="left", vertical="center"))
r += 1
sig_line = Border(bottom=side("C9CCD1"))
for k in ["Name", "Title", "Signature", "Date"]:
    ws.row_dimensions[r].height = 18
    put(f"A{r}", k, font(8, False, "111214"), Alignment(horizontal="left", vertical="bottom"))
    ws.merge_cells(f"B{r}:C{r}"); ws[f"B{r}"].border = sig_line; ws[f"C{r}"].border = sig_line
    put(f"D{r}", k, font(8, False, "111214"), Alignment(horizontal="left", vertical="bottom"))
    ws.merge_cells(f"E{r}:F{r}"); ws[f"E{r}"].border = sig_line; ws[f"F{r}"].border = sig_line
    r += 1
ws.row_dimensions[r].height = 10; r += 1

# Footer
ws.merge_cells(f"A{r}:F{r}"); ws.row_dimensions[r].height = 22; band(r, r, INK)
put(f"A{r}", "Diken Bros Co.  ·  Wadi Saqra, Arar Street 14, Amman 11181, Jordan  ·  Tel 06 416 6660  ·  info@dikenbros.com  ·  dikendelivery.com", font(7.5, False, "FFFFFF"), C)
END = r
ws[f"A{END}"].border = Border(top=side(RED, "medium"))
for c in COLS[1:]: ws[f"{c}{END}"].border = Border(top=side(RED, "medium"))

# Print setup: A4 portrait, one page
ws.print_area = f"A1:F{END}"
ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4, header=0.2, footer=0.2)
ws.print_options.horizontalCentered = True
ws.freeze_panes = None

# Notes sheet
n = wb.create_sheet("How to use")
n.column_dimensions["A"].width = 110
lines = ["How to use the Diken Bros price offer sheet",
 "",
 "1. Save a copy per offer, named after the offer number, e.g. DB-Q-2026-0143.",
 "2. Fill the grey placeholders: offer number, date (Valid until fills itself, 30 days later), client details, your name and mobile, the subject.",
 "3. Type each item: description with Motul part number, quantity, unit, and unit price. The line total, subtotal, discount, net, sales tax and total calculate on their own.",
 "4. Discount and tax percentages sit in column D next to their labels (grey). Leave discount at 0% if none.",
 "5. Payment reads 'As agreed by contract'. Change it only when a deal has no contract.",
 "6. Check the print preview, then export as PDF (File > Save As > PDF, or File > Download > PDF in Google Sheets) before sending to the client.",
 "",
 "Do not change the letterhead, footer or terms without checking with management."]
for i, s in enumerate(lines, 1):
    n[f"A{i}"] = s; n[f"A{i}"].font = font(14 if i == 1 else 10, i == 1); n[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")

out = os.path.join(HERE, "Diken-Price-Offer-Template.xlsx"); wb.save(out); print("written", os.path.basename(out), "rows", END)
